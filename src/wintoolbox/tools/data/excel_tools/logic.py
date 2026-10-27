"""
Lógica pura para manipulação de planilhas Excel (.xlsx, .xlsm) via openpyxl.
Suporta divisão de abas em arquivos individuais, fusão de arquivos em abas ou linhas
e exportação de abas para CSV.

Zero dependências de interface gráfica ou widgets.
"""
import csv
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl


def split_sheets_to_workbooks(xlsx_path: str, output_dir: str) -> List[str]:
    """
    Divide cada aba de uma pasta de trabalho Excel em um arquivo .xlsx independente.
    Retorna a lista de arquivos gerados.
    """
    src_file = Path(xlsx_path)
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    generated_files: List[str] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        new_wb = openpyxl.Workbook()
        new_sheet = new_wb.active
        new_sheet.title = sheet_name

        # Copia valores linha a linha
        for row in sheet.iter_rows(values_only=True):
            new_sheet.append(list(row))

        safe_sheet_name = "".join(c for c in sheet_name if c.isalnum() or c in (" ", "_", "-")).strip()
        out_name = f"{src_file.stem}_{safe_sheet_name or 'Aba'}.xlsx"
        out_path = out_dir / out_name
        new_wb.save(out_path)
        generated_files.append(str(out_path))

    return generated_files


def merge_workbooks_into_sheets(xlsx_paths: List[str], output_path: str) -> str:
    """
    Funde múltiplos arquivos Excel em uma única pasta de trabalho,
    onde cada arquivo de origem se torna uma aba com o nome do arquivo.
    """
    if not xlsx_paths:
        raise ValueError("Nenhum arquivo fornecido para fusão.")

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    target_wb = openpyxl.Workbook()
    # Remove aba padrão inicial
    first_sheet = target_wb.active

    for idx, path in enumerate(xlsx_paths):
        src_p = Path(path)
        src_wb = openpyxl.load_workbook(path, data_only=True)
        sheet_title = src_p.stem[:30]  # Limite de 31 caracteres do Excel

        dest_sheet = target_wb.create_sheet(title=sheet_title)
        active_src_sheet = src_wb.active
        if active_src_sheet:
            for row in active_src_sheet.iter_rows(values_only=True):
                dest_sheet.append(list(row))

    # Remove a aba vazia inicial criada pelo openpyxl
    if first_sheet and len(target_wb.sheetnames) > 1:
        target_wb.remove(first_sheet)

    target_wb.save(out_path)
    return str(out_path)


def merge_workbooks_single_sheet(xlsx_paths: List[str], output_path: str) -> Dict[str, Any]:
    """
    Consolida as linhas de múltiplos arquivos Excel em uma única aba de destino.
    Preserva a primeira linha como cabeçalho do primeiro arquivo e adiciona as linhas seguintes.
    """
    if not xlsx_paths:
        raise ValueError("Nenhum arquivo fornecido para fusão.")

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    target_wb = openpyxl.Workbook()
    dest_sheet = target_wb.active
    dest_sheet.title = "Consolidado"

    total_rows = 0
    header_written = False

    for idx, path in enumerate(xlsx_paths):
        src_wb = openpyxl.load_workbook(path, data_only=True)
        src_sheet = src_wb.active
        if not src_sheet:
            continue

        for row_idx, row in enumerate(src_sheet.iter_rows(values_only=True)):
            if row_idx == 0:
                if not header_written:
                    dest_sheet.append(list(row))
                    header_written = True
                continue

            # Pula linhas completamente vazias
            if all(c is None for c in row):
                continue

            dest_sheet.append(list(row))
            total_rows += 1

    target_wb.save(out_path)
    return {
        "output_path": str(out_path),
        "files_merged": len(xlsx_paths),
        "data_rows_written": total_rows,
    }


def export_sheets_to_csv(
    xlsx_path: str,
    output_dir: str,
    delimiter: str = ";",
    encoding: str = "utf-8",
) -> List[str]:
    """
    Exporta cada aba de um arquivo Excel como um arquivo CSV separado.
    """
    src_file = Path(xlsx_path)
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    generated: List[str] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        safe_sheet_name = "".join(c for c in sheet_name if c.isalnum() or c in (" ", "_", "-")).strip()
        out_name = f"{src_file.stem}_{safe_sheet_name or 'Aba'}.csv"
        out_path = out_dir / out_name

        with open(out_path, "w", encoding=encoding, newline="") as f:
            writer = csv.writer(f, delimiter=delimiter)
            for row in sheet.iter_rows(values_only=True):
                # Substitui None por string vazia
                writer.writerow(["" if c is None else str(c) for c in row])

        generated.append(str(out_path))

    return generated
