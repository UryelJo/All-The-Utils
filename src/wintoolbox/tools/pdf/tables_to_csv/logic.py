"""
Lógica pura para extração de tabelas de documentos PDF e exportação para CSV e Excel (.xlsx).

Zero dependências de interface gráfica ou bibliotecas de widgets.
"""
import csv
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
import pdfplumber


def extract_tables_from_pdf(
    pdf_path: Path,
    pages_to_extract: Optional[List[int]] = None,
) -> List[Dict[str, Any]]:
    """
    Localiza e extrai tabelas de páginas de um arquivo PDF via pdfplumber.
    """
    pdf_path = Path(pdf_path).resolve()
    if not pdf_path.is_file():
        raise FileNotFoundError(f"Arquivo PDF não encontrado: {pdf_path}")

    tables_found: List[Dict[str, Any]] = []

    with pdfplumber.open(pdf_path) as plum:
        for p_idx, page in enumerate(plum.pages):
            p_num = p_idx + 1
            if pages_to_extract and p_num not in pages_to_extract:
                continue

            extracted = page.extract_tables() or []
            for t_idx, raw_table in enumerate(extracted):
                clean_rows: List[List[str]] = []
                max_cols = 0

                for row in raw_table:
                    if not row:
                        continue
                    clean_row = [
                        (str(cell).strip() if cell is not None else "")
                        for cell in row
                    ]
                    clean_rows.append(clean_row)
                    max_cols = max(max_cols, len(clean_row))

                if clean_rows:
                    tables_found.append({
                        "page": p_num,
                        "table_index": t_idx + 1,
                        "rows": clean_rows,
                        "row_count": len(clean_rows),
                        "col_count": max_cols,
                    })

    return tables_found


def export_tables_to_csv(
    tables: List[Dict[str, Any]],
    output_dir: Path,
    delimiter: str = ",",
    encoding: str = "utf-8",
) -> List[Path]:
    """
    Exporta cada tabela extraída para um arquivo CSV individual na pasta especificada.
    """
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_files: List[Path] = []

    for t in tables:
        p_num = t["page"]
        t_num = t["table_index"]
        fn = f"tabela_pag{p_num:02d}_{t_num:02d}.csv"
        file_path = output_dir / fn

        with open(file_path, "w", newline="", encoding=encoding) as f:
            writer = csv.writer(f, delimiter=delimiter)
            for row in t["rows"]:
                writer.writerow(row)

        generated_files.append(file_path)

    return generated_files


def export_tables_to_excel(
    tables: List[Dict[str, Any]],
    output_excel_path: Path,
) -> Path:
    """
    Exporta todas as tabelas extraídas em abas separadas de uma pasta de trabalho do Excel (.xlsx).
    """
    if not tables:
        raise ValueError("Nenhuma tabela fornecida para exportação.")

    output_excel_path = Path(output_excel_path).resolve()
    output_excel_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove a folha padrão inicial vazia
    wb.remove(wb.active)

    for idx, t in enumerate(tables):
        sheet_name = f"Pag{t['page']}_Tab{t['table_index']}"
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel limita em 31 chars

        for r_idx, row in enumerate(t["rows"], start=1):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)

    wb.save(output_excel_path)
    return output_excel_path
