"""
Testes unitários rigorosos para data.excel_tools.logic.
Validação em disco de arquivos reais .xlsx (Black-Box testing) com asserções determinísticas.
"""
import csv
from pathlib import Path
import openpyxl
import pytest
from wintoolbox.tools.data.excel_tools.logic import (
    export_sheets_to_csv,
    merge_workbooks_into_sheets,
    merge_workbooks_single_sheet,
    split_sheets_to_workbooks,
)


def _create_sample_excel(path: Path) -> None:
    wb = openpyxl.Workbook()
    # Aba 1: Vendas
    ws1 = wb.active
    ws1.title = "Vendas"
    ws1.append(["id", "produto", "valor"])
    ws1.append([1, "Monitor", 1200.50])
    ws1.append([2, "Cadeira", 850.00])

    # Aba 2: Custos
    ws2 = wb.create_sheet(title="Custos")
    ws2.append(["categoria", "despesa"])
    ws2.append(["Energia", 450.00])
    ws2.append(["Internet", 150.00])

    wb.save(path)


def test_split_sheets_to_workbooks(tmp_path: Path):
    src = tmp_path / "financeiro.xlsx"
    _create_sample_excel(src)

    out_dir = tmp_path / "split_out"
    generated = split_sheets_to_workbooks(str(src), str(out_dir))

    assert len(generated) == 2

    # Validação do arquivo 1 (Vendas)
    f_vendas = [g for g in generated if "Vendas" in g][0]
    wb_vendas = openpyxl.load_workbook(f_vendas)
    assert wb_vendas.sheetnames == ["Vendas"]
    rows_vendas = list(wb_vendas["Vendas"].iter_rows(values_only=True))
    assert rows_vendas[0] == ("id", "produto", "valor")
    assert rows_vendas[1] == (1, "Monitor", 1200.50)

    # Validação do arquivo 2 (Custos)
    f_custos = [g for g in generated if "Custos" in g][0]
    wb_custos = openpyxl.load_workbook(f_custos)
    assert wb_custos.sheetnames == ["Custos"]
    rows_custos = list(wb_custos["Custos"].iter_rows(values_only=True))
    assert rows_custos[0] == ("categoria", "despesa")
    assert rows_custos[1] == ("Energia", 450.00)


def test_merge_workbooks_into_sheets(tmp_path: Path):
    f1 = tmp_path / "planilha_a.xlsx"
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Dados"
    ws1.append(["A", "B"])
    ws1.append([10, 20])
    wb1.save(f1)

    f2 = tmp_path / "planilha_b.xlsx"
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Dados"
    ws2.append(["C", "D"])
    ws2.append([30, 40])
    wb2.save(f2)

    merged_out = tmp_path / "merged_sheets.xlsx"
    res_path = merge_workbooks_into_sheets([str(f1), str(f2)], str(merged_out))
    assert res_path == str(merged_out)

    wb_res = openpyxl.load_workbook(merged_out)
    assert "planilha_a" in wb_res.sheetnames
    assert "planilha_b" in wb_res.sheetnames

    rows_a = list(wb_res["planilha_a"].iter_rows(values_only=True))
    assert rows_a == [("A", "B"), (10, 20)]

    rows_b = list(wb_res["planilha_b"].iter_rows(values_only=True))
    assert rows_b == [("C", "D"), (30, 40)]


def test_merge_workbooks_single_sheet(tmp_path: Path):
    f1 = tmp_path / "lote1.xlsx"
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["Nome", "Idade"])
    ws1.append(["Carlos", 28])
    wb1.save(f1)

    f2 = tmp_path / "lote2.xlsx"
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["Nome", "Idade"])
    ws2.append(["Diana", 32])
    wb2.save(f2)

    out_file = tmp_path / "lotes_consolidados.xlsx"
    metrics = merge_workbooks_single_sheet([str(f1), str(f2)], str(out_file))

    assert metrics["files_merged"] == 2
    assert metrics["data_rows_written"] == 2

    wb_out = openpyxl.load_workbook(out_file)
    rows = list(wb_out.active.iter_rows(values_only=True))
    assert rows == [
        ("Nome", "Idade"),
        ("Carlos", 28),
        ("Diana", 32),
    ]


def test_export_sheets_to_csv(tmp_path: Path):
    src = tmp_path / "relatorio.xlsx"
    _create_sample_excel(src)

    out_dir = tmp_path / "csv_out"
    csvs = export_sheets_to_csv(str(src), str(out_dir), delimiter=";")
    assert len(csvs) == 2

    csv_vendas = [c for c in csvs if "Vendas" in c][0]
    with open(csv_vendas, "r", encoding="utf-8") as f:
        reader = list(csv.reader(f, delimiter=";"))
        assert reader[0] == ["id", "produto", "valor"]
        assert reader[1] == ["1", "Monitor", "1200.5"]
