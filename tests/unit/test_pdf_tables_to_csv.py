"""
Testes unitários e de validação de efeito real para pdf.tables_to_csv.
"""
import csv
from pathlib import Path
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors

from wintoolbox.tools.pdf.tables_to_csv.logic import (
    extract_tables_from_pdf,
    export_tables_to_csv,
    export_tables_to_excel,
)


def _create_pdf_with_table(path: Path) -> None:
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    data = [
        ["ID", "Produto", "Preco"],
        ["1", "Teclado Mecanico", "250.00"],
        ["2", "Mouse Gamer", "150.00"],
    ]
    t = Table(data)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))
    doc.build([t])


def test_extract_and_export_tables(tmp_path: Path):
    """Testa extração de tabela de PDF e exportação para CSV e Excel."""
    pdf_in = tmp_path / "relatorio_tabela.pdf"
    csv_dir = tmp_path / "saida_csv"
    excel_out = tmp_path / "tabelas.xlsx"

    _create_pdf_with_table(pdf_in)

    # 1. Extração
    tables = extract_tables_from_pdf(pdf_in)
    assert len(tables) >= 1
    first_table = tables[0]
    assert first_table["row_count"] == 3
    assert first_table["col_count"] == 3
    assert first_table["rows"][0] == ["ID", "Produto", "Preco"]
    assert first_table["rows"][1] == ["1", "Teclado Mecanico", "250.00"]

    # 2. Exportação CSV
    csv_files = export_tables_to_csv(tables, csv_dir)
    assert len(csv_files) == len(tables)
    assert csv_files[0].is_file()

    # Validação física do arquivo CSV gerado
    with open(csv_files[0], "r", encoding="utf-8") as f:
        reader = list(csv.reader(f))
        assert reader[0] == ["ID", "Produto", "Preco"]
        assert reader[2] == ["2", "Mouse Gamer", "150.00"]

    # 3. Exportação Excel (.xlsx)
    res_excel = export_tables_to_excel(tables, excel_out)
    assert res_excel.is_file()

    # Validação física do arquivo Excel gerado abrindo com openpyxl
    wb = openpyxl.load_workbook(res_excel)
    sheet = wb.active
    assert sheet.cell(row=1, column=1).value == "ID"
    assert sheet.cell(row=1, column=2).value == "Produto"
    assert sheet.cell(row=2, column=2).value == "Teclado Mecanico"
    assert sheet.cell(row=3, column=3).value == "150.00"
